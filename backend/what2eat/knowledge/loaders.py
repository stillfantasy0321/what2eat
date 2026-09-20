from pathlib import Path
from uuid import UUID
from docx import Document as WordDocument
from langchain_core.documents import Document
from pypdf import PdfReader
from what2eat.errors import AppError

ALLOWED_SUFFIXES = {'.md', '.txt', '.pdf', '.docx'}


def _extract(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in {'.md', '.txt'}:
        try:
            return path.read_text(encoding='utf-8-sig')
        except UnicodeDecodeError as exc:
            raise AppError('invalid_encoding', '文本文件必须使用 UTF-8 编码。', 422) from exc
    if suffix == '.pdf':
        try:
            pages = [(page.extract_text() or '') for page in PdfReader(path).pages]
        except Exception as exc:
            raise AppError('pdf_parse_failed', 'PDF 解析失败。', 422) from exc
        text = '\n\n'.join(pages)
        if not text.strip():
            raise AppError('scanned_pdf', '扫描版 PDF 没有可提取文本，请先进行 OCR。', 422)
        return text
    if suffix == '.docx':
        try:
            document = WordDocument(path)
            return '\n'.join(paragraph.text for paragraph in document.paragraphs)
        except Exception as exc:
            raise AppError('docx_parse_failed', 'DOCX 解析失败。', 422) from exc
    raise AppError('unsupported_file', '不支持该文件格式，仅接受 MD/TXT/PDF/DOCX。', 415)


def load_document(path: Path, document_id: UUID, extracted_limit: int = 1_000_000) -> list[Document]:
    if path.suffix.lower() not in ALLOWED_SUFFIXES:
        raise AppError('unsupported_file', '不支持该文件格式，仅接受 MD/TXT/PDF/DOCX。', 415)
    text = _extract(path).replace('\r\n', '\n').strip()
    if not text:
        raise AppError('empty_document', '文档没有可提取文本。', 422)
    if len(text) > extracted_limit:
        raise AppError('extracted_text_too_large', f'提取文本超过 {extracted_limit} 字符上限。', 413)
    return [Document(page_content=text, metadata={
        'document_id': str(document_id), 'title': path.stem, 'source_path': str(path),
    })]
